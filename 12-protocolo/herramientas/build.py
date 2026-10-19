"""Compila una nota Markdown de la bóveda a PDF con formato APA 7 (TNR 12, doble espacio).

Uso: uv run --with python-docx python build.py ENTRADA.md SALIDA.pdf
Requiere pandoc y libreoffice (soffice) instalados.
"""
import os, re, subprocess, sys, tempfile
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT = "Times New Roman"


def set_font(style, size=12, bold=None, italic=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    style.font.color.rgb = None
    if bold is not None: style.font.bold = bold
    if italic is not None: style.font.italic = italic


def para(style, align=None, indent=None, hanging=None, spacing=None, before=0, after=0):
    spacing = spacing if spacing is not None else (1.0 if os.environ.get("APA_SPACING") == "single" else 2.0)
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE if spacing == 2.0 else WD_LINE_SPACING.SINGLE
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if align is not None: pf.alignment = align
    pf.first_line_indent = Inches(indent) if indent is not None else None
    if hanging:
        pf.left_indent = Inches(hanging); pf.first_line_indent = Inches(-hanging)


def make_reference(path):
    Path(path).write_bytes(subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"],
                                          check=True, capture_output=True).stdout)
    d = Document(path)
    st = d.styles
    def get(name):
        try: return st[name]
        except KeyError:
            from docx.enum.style import WD_STYLE_TYPE
            return st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    for n in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
        s = get(n); set_font(s); para(s, WD_ALIGN_PARAGRAPH.LEFT, indent=0.5 if n in ("Body Text", "First Paragraph") else None)
    para(get("Compact"), WD_ALIGN_PARAGRAPH.LEFT)
    for n, al, b, i in (("Heading 1", WD_ALIGN_PARAGRAPH.CENTER, True, False),
                        ("Heading 2", WD_ALIGN_PARAGRAPH.LEFT, True, False),
                        ("Heading 3", WD_ALIGN_PARAGRAPH.LEFT, True, True)):
        s = get(n); set_font(s, 12, b, i); para(s, al, indent=0); s.paragraph_format.keep_with_next = True
    for n in ("Title", "Subtitle", "Author", "Date"):
        s = get(n); set_font(s, 12, n == "Title"); para(s, WD_ALIGN_PARAGRAPH.CENTER, indent=0)
    s = get("TitlePageTitle"); set_font(s, 12, True); para(s, WD_ALIGN_PARAGRAPH.CENTER, indent=0, before=144)
    s = get("TitlePageText"); set_font(s, 12); para(s, WD_ALIGN_PARAGRAPH.CENTER, indent=0)
    s = get("Bibliography"); set_font(s); para(s, WD_ALIGN_PARAGRAPH.LEFT, hanging=0.5)
    for sec in d.sections:
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
    d.save(path)


def add_page_numbers(doc):
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
        p = sec.header.paragraphs[0] if sec.header.paragraphs else sec.header.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in list(p.runs): r._r.getparent().remove(r._r)
        run = p.add_run()
        for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if t:
                e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
            else:
                e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
            run._r.append(e)
        run.font.name = FONT; run.font.size = Pt(12)


def style_tables(doc):
    for t in doc.tables:
        tblPr = t._tbl.tblPr
        for old in tblPr.findall(qn("w:tblBorders")): tblPr.remove(old)
        b = OxmlElement("w:tblBorders")
        for edge, val in (("top", "single"), ("bottom", "single"), ("left", "nil"), ("right", "nil"), ("insideH", "nil"), ("insideV", "nil")):
            e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), val)
            if val == "single": e.set(qn("w:sz"), "8"); e.set(qn("w:color"), "000000")
            b.append(e)
        tblPr.append(b)
        for old in tblPr.findall(qn("w:tblW")): tblPr.remove(old)
        w = OxmlElement("w:tblW"); w.set(qn("w:type"), "pct"); w.set(qn("w:w"), "5000"); tblPr.append(w)
        ncol = len(t.columns)
        wts = [max(16, min(60, max(len(r.cells[c].text) for r in t.rows))) for c in range(ncol)]
        total = sum(wts); widths = [int(6.5 * 1440 * w / total) for w in wts]
        lay = tblPr.find(qn("w:tblLayout"))
        if lay is None:
            lay = OxmlElement("w:tblLayout"); tblPr.append(lay)
        lay.set(qn("w:type"), "fixed")
        for old in tblPr.findall(qn("w:tblW")): tblPr.remove(old)
        w = OxmlElement("w:tblW"); w.set(qn("w:type"), "dxa"); w.set(qn("w:w"), str(sum(widths))); tblPr.append(w)
        grid = t._tbl.tblGrid
        for gc, wd in zip(grid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(wd))
        for r in t.rows:
            for c, wd in enumerate(widths):
                tcPr = r.cells[c]._tc.get_or_add_tcPr()
                for old in tcPr.findall(qn("w:tcW")): tcPr.remove(old)
                tw = OxmlElement("w:tcW"); tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(wd)); tcPr.insert(0, tw)
        for ri, row in enumerate(t.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    p.paragraph_format.first_line_indent = Inches(0)
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.size = Pt(10); r.font.name = FONT
                        if ri == 0: r.font.bold = True
        if t.rows:
            trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); trPr.append(h)
            for cell in t.rows[0].cells:
                tcPr = cell._tc.get_or_add_tcPr(); bb = OxmlElement("w:tcBorders")
                e = OxmlElement("w:bottom"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); e.set(qn("w:color"), "000000"); bb.append(e); tcPr.append(bb)


def main(src, out):
    text = Path(src).read_text(encoding="utf-8")
    text = re.sub(r"\A---\n.*?\n---\n", "", text, flags=re.S)
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        ref, md, dx = td / "ref.docx", td / "in.md", td / "out.docx"
        make_reference(ref); md.write_text(text, encoding="utf-8")
        subprocess.run(["pandoc", str(md), "-f", "markdown+fenced_divs+raw_attribute", "-o", str(dx),
                        "--reference-doc", str(ref), "-M", "lang=es-MX"], check=True)
        d = Document(dx); style_tables(d); add_page_numbers(d); d.save(dx)
        if os.environ.get("OUT_DOCX"): Path(os.environ["OUT_DOCX"]).write_bytes(dx.read_bytes())
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(td), str(dx)],
                       check=True, capture_output=True)
        Path(out).write_bytes((td / "out.pdf").read_bytes())
    print("OK", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
